import re
import fitz
import pdfplumber
from datetime import datetime as dt
from datetime import date
import pytesseract
from pdf2image import convert_from_path
from PIL import Image
import openpyxl
import pandas as pd
from docx import Document
import logging 

log = logging.getLogger(__name__) 

SUPPORTED_TYPES = ["pdf", "docx", "xlsx"]


def _sanitize_text(text: str) -> str:
    if text is None:
        return ""
    return text.replace("\x00", "").strip()
RESUME_EXTENSIONS = {"pdf", "docx", "xlsx"}

NON_RESUME_PATTERN = re.compile(
    r'\b(invoice|receipt|purchase[\s_-]?order|statement|bill|'
    r'report|brochure|catalogue|catalog|nda|agreement|'
    r'contract|policy|template|screenshot|logo|icon|banner|'
    r'hero|qrcode|apple|google|store|tips|notif|manage|'
    r'microsoft|profile|dummy|avatar)\b',
    re.IGNORECASE
)

RESUME_FILENAME_PATTERNS = re.compile(
    r'\b(resume|cv|curriculum[\s_-]?vitae|cover[\s_-]?letter|'
    r'application|portfolio|biodata|profile)\b',
    re.IGNORECASE
)

SKILLS = [
    "Python", "JavaScript", "TypeScript", "Java", "C", "C++", "C#", "Go", "Rust", "Ruby",
    "PHP", "Perl", "Swift", "Kotlin", "R", "Scala", "Objective-C", "MATLAB", "Shell Scripting",
    "Dart", "Solidity", "GraphQL", "SQL", "NoSQL", "Bash",
    "HTML", "CSS", "SASS", "LESS", "Bootstrap", "Tailwind CSS", "React", "Angular", "Vue.js",
    "Next.js", "Nuxt.js", "Svelte", "jQuery", "Gatsby", "GSAP", "Material UI", "ShadCN",
    "Chakra UI", "Ant Design", "Redux", "Zustand", "Recoil", "MobX", "Storybook", "Vite",
    "Node.js", "Express.js", "Django", "Flask", "FastAPI", "Spring Boot", "ASP.NET",
    "Laravel", "Symfony", "CodeIgniter", "NestJS", "Strapi", "Socket.io", "Hapi.js",
    "PostgreSQL", "MySQL", "MariaDB", "SQLite", "MongoDB", "Cassandra", "Redis", "Elasticsearch",
    "Firebase", "DynamoDB", "CouchDB", "Neo4j", "Snowflake", "BigQuery", "Oracle Database",
    "Microsoft SQL Server", "Supabase", "Prisma", "Sequelize", "Mongoose",
    "AWS", "Azure", "Google Cloud", "IBM Cloud", "Heroku", "DigitalOcean", "Vercel", "Netlify",
    "Docker", "Kubernetes", "Terraform", "Ansible", "Jenkins", "GitLab CI/CD", "GitHub Actions",
    "CircleCI", "Nginx", "Apache", "Prometheus", "Grafana", "Linux Administration",
    "Jest", "Cypress", "Selenium", "Mocha", "Chai", "Puppeteer", "Playwright",
    "React Testing Library", "JUnit", "Pytest", "Postman",
    "React Native", "Flutter", "Android Development", "iOS Development", "Web3.js", "Ethers.js",
    "Machine Learning", "Deep Learning", "Data Science", "Artificial Intelligence",
    "Natural Language Processing", "Computer Vision", "Data Analysis", "Data Visualization",
    "Pandas", "NumPy", "Matplotlib", "Seaborn", "TensorFlow", "Keras", "PyTorch", "Scikit-learn",
    "LangChain", "OpenAI API", "HuggingFace", "Spark", "Hadoop",
    "DSA", "OOP", "REST APIs", "SDLC", "Microservices", "Serverless", "JWT", "OAuth",
    "Agile", "Scrum", "Kanban", "Waterfall", "JIRA", "Confluence", "Git", "GitHub", "GitLab",
    "UI/UX Design", "Figma", "Adobe XD", "Sketch", "Canva", "Photoshop",
    "SEO", "SEM", "Google Analytics", "Accounting", "Financial Analysis", "Technical Writing",
    "Research", "Customer Service", "Sales", "Negotiation", "Public Speaking", "Problem Solving",
    "MS Excel", "MS-PowerPoint", "Team Leadership", "Communication Skills",
]

EMAIL_REGEX = r'\b[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,10}\b'

_LINKEDIN_RAW = (
    r'(?:linkedin\s*[:\-\|]?\s*)?'
    r'(?:https?://)?(?:www\.)?'
    r'linkedin\.com/(?:in|pub|company)/[A-Za-z0-9\-_%]+'
)
_GITHUB_RAW = (
    r'(?:github\s*[:\-\|]?\s*)?'
    r'(?:https?://)?(?:www\.)?'
    r'github\.com/[A-Za-z0-9\-_]+'
)

LINKEDIN_REGEX = re.compile(_LINKEDIN_RAW, re.IGNORECASE)
GITHUB_REGEX = re.compile(_GITHUB_RAW, re.IGNORECASE)

MONTHS_MAP = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}

_EDUCATION_LINE_RE = re.compile(
    r'(?:B\.?Tech|B\.?E\b|B\.?Sc|M\.?Tech|M\.?Sc|MCA|MBA|BCA|BA\b|B\.Com|M\.Com'
    r'|Bachelor|Master|Diploma|Ph\.?D|High\s+School|Secondary|10th|12th'
    r'|University|Institute|College|School)\b',
    re.IGNORECASE,
)

_DATE_RANGE_RE = re.compile(
    r'(?P<start>'
    r'(?:January|February|March|April|May|June|July|August|September|October|November|December|'
    r'Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.?\s+\d{2,4}'
    r'|\d{1,2}[/\-]\d{4}'
    r')'
    r'\s*[-\u2013\u2014to/]+\s*'
    r'(?P<end>'
    r'Present|Current|Till\s+Date|Ongoing|'
    r'(?:January|February|March|April|May|June|July|August|September|October|November|December|'
    r'Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.?\s+\d{2,4}'
    r'|\d{1,2}[/\-]\d{4}'
    r')',
    re.IGNORECASE,
)

_NAME_NOISE = re.compile(
    r'\b(resume|curriculum|vitae|profile|summary|contact|email|phone|address|'
    r'linkedin|github|portfolio|developer|engineer|designer|analyst|manager|'
    r'intern|fresher|objective|declaration|skill|project|education|experience|'
    r'certification|achievement|language|interest|hobby|reference)\b',
    re.IGNORECASE,
)


def _normalise_url(raw: str) -> str:
    raw = raw.strip().rstrip('/')
    if not raw.lower().startswith("http"):
        raw = "https://" + raw
    return raw


def _is_resume_file(file_path: str) -> bool:
    filename = file_path.rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in RESUME_EXTENSIONS:
        return False
    if NON_RESUME_PATTERN.search(filename):
        return False
    return True


def _exp_display(exp_years: float | None) -> str:
    if exp_years is None:
        return "0 mon"
    total_months = round(exp_years * 12)
    if total_months == 0:
        return "0 mon"
    years = total_months // 12
    months = total_months % 12
    if years and months:
        return f"{years} yrs {months} mon"
    elif years:
        return f"{years} yrs"
    return f"{months} mon"


def _parse_month_year(text: str):
    text = text.strip()
    m = re.match(r'(\d{1,2})[/\-](\d{4})', text)
    if m:
        return int(m.group(2)), int(m.group(1))
    m = re.match(r'([A-Za-z]{3,9})[\s\-]+(\d{2,4})', text)
    if m:
        mon_str = m.group(1)[:3].lower()
        year = int(m.group(2))
        if year < 100:
            year += 2000
        mon = MONTHS_MAP.get(mon_str)
        if mon:
            return year, mon
    return None


def extract_hyperlinks_from_pdf(file_obj) -> dict:
    links = {"linkedin": None, "github": None, "portfolio": None, "other": []}
    try:
        if isinstance(file_obj, str):
            doc = fitz.open(file_obj)
        elif hasattr(file_obj, '_path'):
            doc = fitz.open(file_obj._path)
        elif hasattr(file_obj, 'name') and isinstance(file_obj.name, str):
            doc = fitz.open(file_obj.name)
        else:
            if hasattr(file_obj, 'seek'):
                file_obj.seek(0)
            data = file_obj.read()
            doc = fitz.open(stream=data, filetype="pdf")

        for page in doc:
            for link in page.get_links():
                uri = link.get("uri", "")
                if not uri:
                    continue
                if isinstance(uri, bytes):
                    uri = uri.decode("utf-8", errors="ignore")
                uri = uri.strip()
                if not uri or uri.startswith("mailto:") or uri.startswith("tel:"):
                    continue
                if not uri.startswith("http"):
                    uri = "https://" + uri

                if "linkedin.com" in uri and not links["linkedin"]:
                    links["linkedin"] = uri
                elif "github.com" in uri and not links["github"]:
                    links["github"] = uri
                elif not links["portfolio"] and any(
                    kw in uri.lower()
                    for kw in ("portfolio", "netlify", "vercel", "github.io")
                ):
                    links["portfolio"] = uri
                else:
                    links["other"].append(uri)
        doc.close()
    except Exception as e:
        print(f"[WARN] PDF hyperlink extraction failed: {e}")
    return links


def read_pdf(file_path):
    text = ""
    try:
        with pdfplumber.open(file_path) as pdf:
            for page_number, page in enumerate(pdf.pages):
                page_text = page.extract_text()
                if page_text and page_text.strip():
                    text += f"\n--- Page {page_number + 1} ---\n"
                    text += page_text + "\n"

                tables = page.extract_tables()
                if tables:
                    text += f"\n--- Tables on Page {page_number + 1} ---\n"
                    for table in tables:
                        for row in table:
                            clean_row = [cell if cell else "" for cell in row]
                            text += " | ".join(clean_row) + "\n"

        if not text.strip():
            images = convert_from_path(file_path)
            for page_number, image in enumerate(images):
                page_text = pytesseract.image_to_string(image)
                if page_text.strip():
                    text += f"\n--- Page {page_number + 1} (OCR) ---\n"
                    text += page_text + "\n"
        return _sanitize_text(text)
    except Exception as e:
        print(f"Error reading PDF {file_path}: {e}")
        return ""


def read_word(file_path):
    text = ""
    try:
        doc = Document(file_path)
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"
        for table_number, table in enumerate(doc.tables):
            text += f"\n--- Table {table_number + 1} ---\n"
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    text += row_text + "\n"
        return _sanitize_text(text)
    except Exception as e:
        print(f"Error reading Word file {file_path}: {e}")
        return ""


def read_excel(file_path):
    text = ""
    try:
        workbook = openpyxl.load_workbook(file_path)
        sheet_names = workbook.sheetnames
        for sheet_name in sheet_names:
            text += f"\n--- Sheet: {sheet_name} ---\n"
            df = pd.read_excel(file_path, sheet_name=sheet_name)
            df = df.dropna(how="all")
            df = df.dropna(axis=1, how="all")
            text += df.to_string(index=False)
            text += "\n"
        return _sanitize_text(text)
    except Exception as e:
        print(f"Error reading Excel file {file_path}: {e}")
        return ""


def get_excel_summary(file_path):
    summary = {}
    try:
        workbook = openpyxl.load_workbook(file_path)
        summary["total_sheets"] = len(workbook.sheetnames)
        summary["sheet_names"] = workbook.sheetnames
        summary["sheets_detail"] = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            summary["sheets_detail"].append({
                "name": sheet_name,
                "total_rows": sheet.max_row,
                "total_columns": sheet.max_column,
            })
        return summary
    except Exception as e:
        print(f"Error getting Excel summary: {e}")
        return {}


def read_attachment(file_path: str) -> str:
    ext = file_path.split(".")[-1].lower()
    if ext not in SUPPORTED_TYPES:
        return ""
    if ext == "pdf":
        return read_pdf(file_path)
    elif ext == "docx":
        return read_word(file_path)
    elif ext == "xlsx":
        return read_excel(file_path)
    return ""


def extract_skills(text: str) -> list:
    lower_text = text.lower()
    found = []
    seen = set()
    for skill in SKILLS:
        if re.search(rf'\b{re.escape(skill.lower())}\b', lower_text):
            if skill.lower() not in seen:
                found.append(skill)
                seen.add(skill.lower())
    return found


def extract_name_from_text(text: str) -> str | None:
    raw_lines = [l.strip() for l in text.splitlines() if l.strip()][:25]
    joined_lines = []
    i = 0
    while i < len(raw_lines):
        line = raw_lines[i]
        if (
            i + 1 < len(raw_lines)
            and re.match(r'^[A-Z]$', line)
            and re.match(r'^[A-Z][A-Z\s]+$', raw_lines[i + 1])
        ):
            joined_lines.append(line + raw_lines[i + 1])
            i += 2
        else:
            joined_lines.append(line)
            i += 1

    for line in joined_lines:
        line = re.split(r'\s*[–—\-]\s*', line)[0].strip()

        if re.search(r'[@:/\d]', line):
            continue
        if len(line) > 50:
            continue
        if _NAME_NOISE.search(line):
            continue
        if re.search(r'[|,;•·]', line):
            continue

        words = line.split()
        if not (2 <= len(words) <= 4):
            continue

        if all(re.match(r'^[A-Z][a-zA-Z\'-]*$', w) for w in words):
            return " ".join(w.title() for w in words)

        if all(re.match(r'^[A-Z]+$', w) for w in words):
            return " ".join(w.title() for w in words)

    return None

def _strip_education_lines(text: str) -> str:
    cleaned = []
    for line in text.splitlines():
        if _EDUCATION_LINE_RE.search(line):
            cleaned.append("")
        else:
            cleaned.append(line)
    return "\n".join(cleaned)


def _fix_2digit_year(raw: str) -> str:
    return re.sub(
        r'^([A-Za-z]{3,9})\s+(\d{2})$',
        lambda m: f"{m.group(1)} {2000 + int(m.group(2))}",
        raw.strip(),
    )


def _normalize_date_fixed(d: str):
    if not d:
        return None
    d = d.strip()
    if d.lower() in ("present", "current", "till date", "ongoing"):
        return None
    d = _fix_2digit_year(d)
    m = re.match(r'(\d{1,2})[/\-](\d{4})', d)
    if m:
        return int(m.group(2)), int(m.group(1))
    m = re.match(r'([A-Za-z]{3,9})[\s\-]+(\d{4})', d)
    if m:
        mon = MONTHS_MAP.get(m.group(1)[:3].lower())
        if mon:
            return int(m.group(2)), mon
    try:
        from dateutil import parser as date_parser
        parsed = date_parser.parse(d, default=dt(2000, 1, 1))
        if 2000 <= parsed.year <= 2035:
            return (parsed.year, parsed.month)
    except Exception:
        pass
    return None


def _scan_date_ranges(text: str) -> int:
    today = date.today()
    seen  = set()
    ranges = []

    for m in _DATE_RANGE_RE.finditer(text):
        start_str = m.group("start").strip()
        end_str   = m.group("end").strip()

        start = _normalize_date_fixed(start_str)
        if not start:
            continue

        if re.match(r'present|current|till\s+date|ongoing', end_str, re.IGNORECASE):
            end = (today.year, today.month)
        else:
            end = _normalize_date_fixed(end_str)
            if not end:
                continue
            if date(end[0], end[1], 1) > today:
                continue

        key = (start, end)
        if key in seen:
            continue
        seen.add(key)

        s = start[0] * 12 + start[1]
        e = end[0]   * 12 + end[1]
        if 0 < e - s < 600:
            ranges.append((s, e))

    if not ranges:
        return 0

    # Merge overlapping ranges — prevents double-counting same-employer projects
    ranges.sort()
    merged = [ranges[0]]
    for s, e in ranges[1:]:
        if s <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], e))
        else:
            merged.append((s, e))

    return sum(e - s for s, e in merged)


def extract_experience(text: str) -> float | None:
    def _to_years(total_months: int) -> float:
        return round(total_months / 12, 2)

    match = re.search(
        r'(\d+)\+?\s*(?:years?|yrs?)\s*(?:and\s+)?(\d+)\s*(?:months?|mos?)',
        text, re.IGNORECASE,
    )
    if match:
        return _to_years(int(match.group(1)) * 12 + int(match.group(2)))

    match = re.search(
        r'(\d+\.\d+)\+?\s*(?:years?\'?|yrs?\'?)[\s\w]{0,20}?(?:of\s+)?(?:experience|internship)?',
        text, re.IGNORECASE,
    )
    if match:
        return _to_years(round(float(match.group(1)) * 12))

    match = re.search(
        r'(\d+)\+?\s*(?:years?\'?|yrs?\'?)\W{0,5}(?:of\s+)?(?:experience|internship)',
        text, re.IGNORECASE,
    )
    if match:
        return _to_years(int(match.group(1)) * 12)

    match = re.search(
        r'(\d+)\+?\s*(?:months?|mos?)[\s\w]{0,20}?(?:of\s+)?(?:experience|internship)',
        text, re.IGNORECASE,
    )
    if match:
        return _to_years(int(match.group(1)))

    match = re.search(r'(\d+)\s*[-\u2013]?\s*month\s+internship', text, re.IGNORECASE)
    if match:
        return _to_years(int(match.group(1)))

    match = re.search(r'interned?\s+(?:for\s+)?(\d+)\s*(?:months?|mos?)', text, re.IGNORECASE)
    if match:
        return _to_years(int(match.group(1)))

    for line in text.splitlines():
        if re.search(r'\bintern\b', line, re.IGNORECASE):
            m = re.search(r'(\d+)\s*(?:months?|mos?)', line, re.IGNORECASE)
            if m:
                return _to_years(int(m.group(1)))

    work_text = _strip_education_lines(text)
    total_months = _scan_date_ranges(work_text)
    if total_months > 0:
        return _to_years(total_months)

    return None


def extract_email(text: str):
    emails = re.findall(EMAIL_REGEX, text)
    return emails[0] if emails else None


def extract_profile_urls(text: str) -> dict:
    linkedin = github = None
    m = LINKEDIN_REGEX.search(text)
    if m:
        raw = re.sub(r'^(?:linkedin\s*[:\-\|]\s*)', '', m.group(0), flags=re.IGNORECASE).strip()
        linkedin = _normalise_url(raw)
    m = GITHUB_REGEX.search(text)
    if m:
        raw = re.sub(r'^(?:github\s*[:\-\|]\s*)', '', m.group(0), flags=re.IGNORECASE).strip()
        github = _normalise_url(raw)
    return {"linkedin": linkedin, "github": github}


def extract_from_attachment_text(text: str, file_path: str = None) -> dict:
    if not text:
        return {
            "skills": [],
            "experience": "0 mon",
            "email": None,
            "linkedin": None,
            "github": None,
            "portfolio": None,
            "name": None,
        }

    profile_urls = extract_profile_urls(text)
    linkedin = profile_urls["linkedin"]
    github = profile_urls["github"]
    portfolio = None

    if file_path and file_path.lower().endswith(".pdf"):
        pdf_links = extract_hyperlinks_from_pdf(file_path)
        linkedin = pdf_links.get("linkedin") or linkedin
        github = pdf_links.get("github") or github
        portfolio = pdf_links.get("portfolio") or portfolio

    return {
        "skills": extract_skills(text),
        "experience": _exp_display(extract_experience(text)),
        "email": extract_email(text),
        "linkedin": linkedin,
        "github": github,
        "portfolio": portfolio,
        "name": extract_name_from_text(text),
    }


def process_attachment(file_path: str) -> dict:
    if not _is_resume_file(file_path):
        return {
            "skills": [],
            "experience": "0 mon",
            "email": None,
            "linkedin": None,
            "github": None,
            "portfolio": None,
            "name": None,
            "raw_text": None,
        }
    raw_text = read_attachment(file_path)
    extracted = extract_from_attachment_text(raw_text, file_path=file_path)
    extracted["raw_text"] = raw_text
    return extracted

async def process_and_analyze_attachment(
    file_path: str,
    provider: str,
    db,
    candidate_name: str = "Candidate",
) -> dict:
    raw_text = read_attachment(file_path)
    extracted = extract_from_attachment_text(raw_text, file_path=file_path)
    exp_years = extract_experience(raw_text)

    resume_name = extracted.get("name") or None
    resume_email = extracted.get("email") or None

    if not resume_name:
        log.warning(
            "[PARSE] Could not extract name from resume %s — "
            "falling back to email sender: %s", file_path, candidate_name
        )
        resume_name = candidate_name

    parsed_resume = {
        "name": resume_name,
        "email": resume_email or "",   
        "linkedin": extracted.get("linkedin") or "",
        "github": extracted.get("github") or "",
        "portfolio": extracted.get("portfolio") or "",
        "skills": extracted.get("skills", []),
        "exp_years": exp_years,
        "experience": [],
        "work_experiences": [],
        "education": [],
        "projects": [],
        "certifications": [],
        "raw_text": raw_text,
    }

    from resume_analyzer.chains.resume_chain import run_chain
    result = await run_chain(parsed_resume)

    if resume_email:
        result["email"] = resume_email
    result["name"] = resume_name
    return result


def is_relevant_attachment(filename: str) -> bool:
    if not filename:
        return False
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in RESUME_EXTENSIONS:
        return False
    NON_RESUME_PATTERNS = re.compile(
        r'\b(invoice|receipt|purchase[\s_-]?order|statement|bill|'
        r'report|presentation|brochure|catalogue|catalog|nda|agreement|'
        r'contract|policy|template|sample|test[\s_-]?case|screenshot)\b',
        re.IGNORECASE
    )
    if NON_RESUME_PATTERNS.search(filename):
        return False
    return True


def filter_job_attachments(
    attachment_names: list[str],
    subject: str,
    body: str,
    sender_email: str = "",
) -> list[str]:
    from services.extractor import is_job_application
    if not is_job_application(subject, body, sender_email):
        return []
    return [name for name in attachment_names if is_relevant_attachment(name)]